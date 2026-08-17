"""Simulation study for the finite-interval/exact-point support model.

All performance scenarios use at least ten independent study blocks.  The core
scenarios use a reproducible ten-study template containing each of the eight
empirical block configurations once plus two resampled blocks; this preserves both
positive-width and exact point-versus-point supports.  A generalizability
experiment resamples empirical study-block configurations at K=10, 20, and 30.
Simulation summaries report bias, RMSE, and confidence-interval calibration
for both normal-Wald and study-count Student-t intervals.
"""

from __future__ import annotations

import argparse
import json
import math
import time
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.special import expit, log_ndtr, logit
from scipy.stats import gamma, t as student_t

from interval_drma_core import (
    FastReml,
    Row,
    Study,
    Support,
    parametric_bootstrap_interval,
)


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_BOOK = ROOT / "data" / "analytic_contrasts.xlsx"
OUTPUT = Path(__file__).with_name("mixed_support_simulation_results.json")

SCHULTZ_REFID = 32508669
REGION_PARAMS = {
    "Netherlands": (2.640331, 0.220433),
    "NC": (1.809830, 5.193423),
    "PA": (2.059092, 2.108859),
    "Wisconsin": (1.900381, 6.181165),
    "Iowa": (2.48, 1.44),
}

BASE_BETA = -0.069469
BASE_TAU = 0.129253
RHO_WORKING = 0.5


def number(value):
    if isinstance(value, str) and value.strip().lower() in {"inf", ">inf"}:
        return math.inf
    return float(value)


def read_rows(book: Path) -> list[Row]:
    if not book.is_file():
        raise FileNotFoundError(
            f"Analytic workbook not found: {book}. "
            "Supply its location with --data-file."
        )
    workbook = load_workbook(book, read_only=True, data_only=True)
    sheet_name = (
        "Analytic contrasts"
        if "Analytic contrasts" in workbook.sheetnames
        else "Selected Obs"
    )
    sheet = workbook[sheet_name]
    values = list(
        sheet.iter_rows(min_row=1, max_row=34, min_col=1, max_col=20, values_only=True)
    )
    headers = values[0]
    records = [dict(zip(headers, row)) for row in values[1:] if row[0] is not None]
    rows = []

    for record in records:
        refid = int(record["Refid"])
        region = str(record["Region"]).strip()
        c_lower = number(record["Contrast_distance_Lower"])
        c_upper = number(record["Contrast_distance_Upper"])
        k_lower = number(record["Control_distance_Lower"])
        k_upper = number(record["Control_distance_Upper"])
        dmax = None if record["d_max"] is None else number(record["d_max"])

        if refid == SCHULTZ_REFID:
            comparison = Support(c_upper, c_upper, point=True)
            reference = Support(k_upper, k_upper, point=True)
        else:
            c_from_dmax = math.isinf(c_upper)
            if c_from_dmax:
                if dmax is None:
                    raise ValueError("Finite d_max is required for the application")
                c_upper = dmax
            comparison = Support(c_lower, c_upper, upper_from_dmax=c_from_dmax)

            if int(record["Type"]) == 1:
                k_from_dmax = math.isinf(k_upper)
                if k_from_dmax:
                    if dmax is None:
                        raise ValueError("Finite d_max is required for the application")
                    k_upper = dmax
                reference = Support(k_lower, k_upper, upper_from_dmax=k_from_dmax)
            else:
                if dmax is None:
                    raise ValueError("Finite d_max is required for binary buffer contrasts")
                reference = Support(c_upper, dmax, upper_from_dmax=True)

        se = (math.log(float(record["UL"])) - math.log(float(record["LL"]))) / 3.92
        log_effect = math.log(float(record["EM"]))
        effect_measure = str(record["effect_measure"]).strip().upper()
        rows.append(
            Row(
                refid=refid,
                region=region,
                comparison=comparison,
                reference=reference,
                se=se,
                log_effect=log_effect,
                effect_measure=effect_measure,
                author=str(record.get("Author") or "").strip(),
                comparison_label=str(
                    record.get("trt1_name") or record.get("value") or ""
                ).strip(),
                reference_label=str(record.get("trt2_name") or "").strip(),
                dmax=dmax,
                dmax_reason=str(record.get("d_max Reason") or "").strip(),
                mapping_type=int(record["Type"]),
            )
        )
    return rows


def build_studies(rows: list[Row], rho: float) -> list[Study]:
    studies = []
    for refid in dict.fromkeys(row.refid for row in rows):
        study_rows = [row for row in rows if row.refid == refid]
        se = np.asarray([row.se for row in study_rows])
        corr = np.full((len(study_rows), len(study_rows)), rho)
        np.fill_diagonal(corr, 1.0)
        sigma = np.diag(se) @ corr @ np.diag(se)
        studies.append(Study(refid, study_rows, sigma))
    return studies


def gamma_mean(support: Support, shape: float, scale: float) -> float:
    if support.point:
        return support.lower

    def log_difference(log_large: float, log_small: float) -> float:
        if log_small == -math.inf:
            return log_large
        delta = log_small - log_large
        if delta >= 0:
            raise FloatingPointError("Interval probability is numerically unresolved")
        return log_large + math.log(-math.expm1(delta))

    def log_interval_probability(parameter_shape: float) -> float:
        if math.isinf(support.upper):
            return float(gamma.logsf(support.lower, a=parameter_shape, scale=scale))
        if support.lower <= 0:
            return float(gamma.logcdf(support.upper, a=parameter_shape, scale=scale))
        log_cdf_upper = float(
            gamma.logcdf(support.upper, a=parameter_shape, scale=scale)
        )
        if log_cdf_upper <= math.log(0.5):
            return log_difference(
                log_cdf_upper,
                float(gamma.logcdf(support.lower, a=parameter_shape, scale=scale)),
            )
        return log_difference(
            float(gamma.logsf(support.lower, a=parameter_shape, scale=scale)),
            float(gamma.logsf(support.upper, a=parameter_shape, scale=scale)),
        )

    log_denom = log_interval_probability(shape)
    log_numer_probability = log_interval_probability(shape + 1.0)
    if not math.isfinite(log_denom) or not math.isfinite(log_numer_probability):
        raise FloatingPointError(
            "Gamma support probability is too small to resolve numerically"
        )
    result = shape * scale * math.exp(log_numer_probability - log_denom)
    tolerance = 1e-10 * max(1.0, abs(result))
    if result < support.lower - tolerance or (
        math.isfinite(support.upper) and result > support.upper + tolerance
    ):
        raise FloatingPointError("Computed Gamma conditional mean lies outside its support")
    return result


def lognormal_mean(support: Support, shape: float, scale: float) -> float:
    if support.point:
        return support.lower
    arithmetic_mean = shape * scale
    sigma2 = math.log1p(1.0 / shape)
    sigma = math.sqrt(sigma2)
    mu = math.log(arithmetic_mean) - sigma2 / 2.0

    def z(value: float, shifted: bool = False) -> float:
        if value <= 0:
            return -math.inf
        if math.isinf(value):
            return math.inf
        return (math.log(value) - mu - (sigma2 if shifted else 0.0)) / sigma

    def log_difference(log_large: float, log_small: float) -> float:
        if log_small == -math.inf:
            return log_large
        delta = log_small - log_large
        if delta >= 0:
            raise FloatingPointError("Interval probability is numerically unresolved")
        return log_large + math.log(-math.expm1(delta))

    def normal_log_interval(lower_z: float, upper_z: float) -> float:
        if upper_z == math.inf:
            return float(log_ndtr(-lower_z))
        if lower_z == -math.inf:
            return float(log_ndtr(upper_z))
        log_cdf_upper = float(log_ndtr(upper_z))
        if log_cdf_upper <= math.log(0.5):
            return log_difference(log_cdf_upper, float(log_ndtr(lower_z)))
        return log_difference(float(log_ndtr(-lower_z)), float(log_ndtr(-upper_z)))

    log_denom = normal_log_interval(z(support.lower), z(support.upper))
    log_numer_probability = normal_log_interval(
        z(support.lower, shifted=True), z(support.upper, shifted=True)
    )
    if not math.isfinite(log_denom) or not math.isfinite(log_numer_probability):
        raise FloatingPointError(
            "Lognormal support probability is too small to resolve numerically"
        )
    result = arithmetic_mean * math.exp(log_numer_probability - log_denom)
    tolerance = 1e-10 * max(1.0, abs(result))
    if result < support.lower - tolerance or (
        math.isfinite(support.upper) and result > support.upper + tolerance
    ):
        raise FloatingPointError(
            "Computed Lognormal conditional mean lies outside its support"
        )
    return result


def midpoint(support: Support) -> float:
    if support.point:
        return support.lower
    return (support.lower + support.upper) / 2.0


def design_vector(
    study: Study,
    distribution: str,
    dmax_multiplier: float = 1.0,
    region_params: dict[str, tuple[float, float]] | None = None,
    infinite_open_tails: bool = False,
) -> np.ndarray:
    if region_params is None:
        region_params = REGION_PARAMS
    values = []
    for row in study.rows:
        shape, scale = region_params.get(row.region, region_params["Iowa"])
        a = row.comparison.with_dmax_multiplier(dmax_multiplier)
        b = row.reference.with_dmax_multiplier(dmax_multiplier)
        if infinite_open_tails:
            a = (
                Support(a.lower, math.inf)
                if a.upper_from_dmax and not a.point
                else a
            )
            b = (
                Support(b.lower, math.inf)
                if b.upper_from_dmax and not b.point
                else b
            )
            if distribution in {"uniform", "midpoint"}:
                raise ValueError("Midpoints are undefined for right-unbounded supports")
        if distribution == "gamma":
            mean_fn = lambda support: gamma_mean(support, shape, scale)
        elif distribution == "lognormal":
            mean_fn = lambda support: lognormal_mean(support, shape, scale)
        elif distribution in {"uniform", "midpoint"}:
            mean_fn = midpoint
        else:
            raise ValueError(f"Unknown distribution: {distribution}")
        values.append(mean_fn(a) - mean_fn(b))
    return np.asarray(values)


def logistic_support_probability(
    support: Support,
    shape: float,
    scale: float,
    intercept: float,
    slopes: np.ndarray,
    curvature: float = 0.0,
    nodes: int = 32,
) -> np.ndarray:
    if support.point:
        return expit(
            intercept
            + slopes * support.lower
            + curvature * math.log1p(support.lower)
        )

    legendre_x, legendre_w = np.polynomial.legendre.leggauss(nodes)
    distance = (support.upper - support.lower) * (legendre_x + 1.0) / 2.0 + support.lower
    weights = legendre_w * (support.upper - support.lower) / 2.0
    density = gamma.pdf(distance, a=shape, scale=scale)
    conditional_weights = weights * density
    conditional_weights /= conditional_weights.sum()
    probabilities = expit(
        intercept
        + slopes[:, None] * distance[None, :]
        + curvature * np.log1p(distance)[None, :]
    )
    return probabilities @ conditional_weights


def simulate_linear(
    rng: np.random.Generator,
    studies_true: list[Study],
    studies_est: list[Study],
    x_true: list[np.ndarray],
    x_est: list[np.ndarray],
    beta: float,
    tau: float,
    replicates: int,
) -> dict:
    fitter = FastReml(studies_est, x_est)
    chol = [np.linalg.cholesky(study.sigma) for study in studies_true]
    beta_est = np.empty(replicates)
    tau_est = np.empty(replicates)
    se_beta = np.empty(replicates)

    for r in range(replicates):
        random_slopes = rng.normal(0.0, tau, size=len(studies_true))
        y_list = []
        for k, (x, factor) in enumerate(zip(x_true, chol)):
            error = factor @ rng.normal(size=len(x))
            y_list.append((beta + random_slopes[k]) * x + error)
        beta_est[r], tau_est[r] = fitter.fit(y_list)
        se_beta[r] = fitter.standard_error(tau_est[r])

    return summarize(
        beta_est,
        tau_est,
        beta,
        tau,
        se_beta=se_beta,
        n_studies=len(studies_est),
    )


def simulate_logistic_aggregation(
    rng: np.random.Generator,
    studies: list[Study],
    x_est: list[np.ndarray],
    beta: float,
    tau: float,
    replicates: int,
    intercept: float = -2.2,
    curvature: float = 0.0,
) -> dict:
    fitter = FastReml(studies, x_est)
    chol = [np.linalg.cholesky(study.sigma) for study in studies]
    random_slopes = rng.normal(0.0, tau, size=(replicates, len(studies)))
    errors = [rng.normal(size=(replicates, len(study.rows))) @ factor.T for study, factor in zip(studies, chol)]

    true_y = []
    for k, study in enumerate(studies):
        slopes = beta + random_slopes[:, k]
        values = np.empty((replicates, len(study.rows)))
        for j, row in enumerate(study.rows):
            shape, scale = REGION_PARAMS.get(row.region, REGION_PARAMS["Iowa"])
            p_a = logistic_support_probability(
                row.comparison, shape, scale, intercept, slopes, curvature
            )
            p_b = logistic_support_probability(
                row.reference, shape, scale, intercept, slopes, curvature
            )
            values[:, j] = logit(np.clip(p_a, 1e-12, 1 - 1e-12)) - logit(
                np.clip(p_b, 1e-12, 1 - 1e-12)
            )
        true_y.append(values)

    beta_est = np.empty(replicates)
    tau_est = np.empty(replicates)
    se_beta = np.empty(replicates)
    for r in range(replicates):
        y_list = [values[r] + error[r] for values, error in zip(true_y, errors)]
        beta_est[r], tau_est[r] = fitter.fit(y_list)
        se_beta[r] = fitter.standard_error(tau_est[r])
    return summarize(
        beta_est,
        tau_est,
        beta,
        tau,
        se_beta=se_beta,
        n_studies=len(studies),
    )


def perturbed_region_params(
    rng: np.random.Generator,
    coefficient_of_variation: float,
) -> dict[str, tuple[float, float]]:
    """Draw unbiased positive perturbations of each regional Gamma parameter.

    The scenario is intentionally expressed through a coefficient of variation
    because standard errors for every external regional fit were not available.
    Shape and scale are perturbed independently on the log scale.
    """
    sigma2 = math.log1p(coefficient_of_variation**2)
    sigma = math.sqrt(sigma2)
    parameters = {}
    for region, (shape, scale) in REGION_PARAMS.items():
        shape_draw = rng.lognormal(math.log(shape) - sigma2 / 2.0, sigma)
        scale_draw = rng.lognormal(math.log(scale) - sigma2 / 2.0, sigma)
        parameters[region] = (float(shape_draw), float(scale_draw))
    return parameters


def simulate_parameter_uncertainty(
    rng: np.random.Generator,
    studies: list[Study],
    x_true: list[np.ndarray],
    beta: float,
    tau: float,
    replicates: int,
    coefficient_of_variation: float,
) -> dict:
    """Evaluate plug-in estimation with uncertain exposure parameters."""
    chol = [np.linalg.cholesky(study.sigma) for study in studies]
    beta_est = np.empty(replicates)
    tau_est = np.empty(replicates)
    se_beta = np.empty(replicates)

    for r in range(replicates):
        fitted_parameters = perturbed_region_params(
            rng, coefficient_of_variation
        )
        x_est = [
            design_vector(study, "gamma", region_params=fitted_parameters)
            for study in studies
        ]
        fitter = FastReml(studies, x_est)
        random_slopes = rng.normal(0.0, tau, size=len(studies))
        y_list = []
        for k, (x, factor) in enumerate(zip(x_true, chol)):
            error = factor @ rng.normal(size=len(x))
            y_list.append((beta + random_slopes[k]) * x + error)
        beta_est[r], tau_est[r] = fitter.fit(y_list)
        se_beta[r] = fitter.standard_error(tau_est[r])

    summary = summarize(
        beta_est,
        tau_est,
        beta,
        tau,
        se_beta=se_beta,
        n_studies=len(studies),
    )
    summary["parameter_cv"] = coefficient_of_variation
    return summary


def simulate_resampled_study_count(
    rng: np.random.Generator,
    base_studies: list[Study],
    base_x: list[np.ndarray],
    beta: float,
    tau: float,
    replicates: int,
    n_studies: int,
    templates: int = 25,
) -> dict:
    """Evaluate generalizability across K using resampled empirical blocks.

    Study blocks are sampled with replacement from the eight observed block
    configurations.  A fixed set of independently sampled templates is used within each
    K cell so that the experiment varies study count and support configuration
    without treating one arbitrarily selected subset as representative.
    """
    if n_studies < 2:
        raise ValueError("At least two study blocks are required")
    if templates < 1 or templates > replicates:
        raise ValueError("templates must be between one and replicates")

    beta_est = np.empty(replicates)
    tau_est = np.empty(replicates)
    se_beta = np.empty(replicates)
    template_sizes = np.full(templates, replicates // templates, dtype=int)
    template_sizes[: replicates % templates] += 1
    cursor = 0

    for template_replicates in template_sizes:
        indices = rng.choice(len(base_studies), size=n_studies, replace=True)
        studies = [base_studies[index] for index in indices]
        x = [base_x[index] for index in indices]
        fitter = FastReml(studies, x)
        chol = [np.linalg.cholesky(study.sigma) for study in studies]
        for _ in range(int(template_replicates)):
            random_slopes = rng.normal(0.0, tau, size=n_studies)
            y_list = [
                (beta + random_slopes[k]) * x[k]
                + chol[k] @ rng.normal(size=len(x[k]))
                for k in range(n_studies)
            ]
            beta_r, tau_r = fitter.fit(y_list)
            beta_est[cursor] = beta_r
            tau_est[cursor] = tau_r
            se_beta[cursor] = fitter.standard_error(tau_r)
            cursor += 1

    summary = summarize(
        beta_est,
        tau_est,
        beta,
        tau,
        se_beta=se_beta,
        n_studies=n_studies,
    )
    summary.update(
        {
            "n_studies": n_studies,
            "templates": templates,
            "sampling_scheme": (
                "Empirical study-block configurations sampled with replacement"
            ),
        }
    )
    return summary


def ten_study_template_indices(n_base_studies: int, seed: int = 8675309) -> list[int]:
    """Return a reproducible K=10 template retaining every empirical block configuration."""
    if n_base_studies != 8:
        raise ValueError(
            "The empirical template is expected to contain eight block configurations"
        )
    rng = np.random.default_rng(seed)
    indices = list(range(n_base_studies))
    indices.extend(int(value) for value in rng.choice(n_base_studies, size=2, replace=True))
    rng.shuffle(indices)
    return indices


def summarize(
    beta_est,
    tau_est,
    beta_true,
    tau_true,
    se_beta=None,
    n_studies=None,
) -> dict:
    beta_error = beta_est - beta_true
    tau_error = tau_est - tau_true
    result = {
        "replicates": int(len(beta_est)),
        "beta_true": beta_true,
        "tau_true": tau_true,
        "mean_beta": float(beta_est.mean()),
        "mean_tau": float(tau_est.mean()),
        "beta_bias": float(beta_error.mean()),
        "beta_rmse": float(np.sqrt(np.mean(beta_error**2))),
        "tau_bias": float(tau_error.mean()),
        "tau_rmse": float(np.sqrt(np.mean(tau_error**2))),
        "tau_zero_boundary_rate": float(np.mean(np.asarray(tau_est) == 0.0)),
        "convergence_rate": 1.0,
    }
    if se_beta is not None:
        if n_studies is None or n_studies < 2:
            raise ValueError("At least two studies are required for study-t inference")
        se_beta = np.asarray(se_beta, dtype=float)
        normal_critical = 1.959963984540054
        study_df = n_studies - 1
        study_t_critical = float(student_t.ppf(0.975, study_df))
        normal_half_width = normal_critical * se_beta
        study_t_half_width = study_t_critical * se_beta
        normal_covered = np.abs(beta_error) <= normal_half_width
        study_t_covered = np.abs(beta_error) <= study_t_half_width
        normal_coverage = float(normal_covered.mean())
        study_t_coverage = float(study_t_covered.mean())
        replicates = len(beta_error)
        result["inference"] = {
            "study_t_df": study_df,
            "normal_coverage": normal_coverage,
            "study_t_coverage": study_t_coverage,
            "normal_coverage_mcse": float(
                math.sqrt(normal_coverage * (1.0 - normal_coverage) / replicates)
            ),
            "study_t_coverage_mcse": float(
                math.sqrt(study_t_coverage * (1.0 - study_t_coverage) / replicates)
            ),
            "normal_mean_width": float((2.0 * normal_half_width).mean()),
            "study_t_mean_width": float((2.0 * study_t_half_width).mean()),
        }
        if beta_true == 0:
            normal_type1 = 1.0 - normal_coverage
            study_t_type1 = 1.0 - study_t_coverage
            result["inference"].update(
                {
                    "normal_type1_error": float(normal_type1),
                    "study_t_type1_error": float(study_t_type1),
                    "normal_type1_error_mcse": float(
                        math.sqrt(normal_type1 * (1.0 - normal_type1) / replicates)
                    ),
                    "study_t_type1_error_mcse": float(
                        math.sqrt(study_t_type1 * (1.0 - study_t_type1) / replicates)
                    ),
                }
            )
    return result


def simulate_fixed(
    rng: np.random.Generator,
    studies: list[Study],
    x: list[np.ndarray],
    beta: float,
    replicates: int,
) -> dict:
    inv = [np.linalg.inv(study.sigma) for study in studies]
    den = sum(float(xk @ vk @ xk) for xk, vk in zip(x, inv))
    chol = [np.linalg.cholesky(study.sigma) for study in studies]
    estimates = np.empty(replicates)
    for r in range(replicates):
        num = 0.0
        for xk, vk, factor in zip(x, inv, chol):
            y = beta * xk + factor @ rng.normal(size=len(xk))
            num += float(xk @ vk @ y)
        estimates[r] = num / den
    error = estimates - beta
    se_beta = math.sqrt(1.0 / den)
    covered = np.abs(error) <= 1.959963984540054 * se_beta
    return {
        "beta_true": beta,
        "mean_beta": float(estimates.mean()),
        "beta_bias": float(error.mean()),
        "beta_rmse": float(np.sqrt(np.mean(error**2))),
        "normal_coverage": float(covered.mean()),
        "normal_mean_width": 2.0 * 1.959963984540054 * se_beta,
    }


def main(replicates: int, data_file: Path) -> dict:
    rows = read_rows(data_file)
    empirical_rho05 = build_studies(rows, RHO_WORKING)
    empirical_rho03 = build_studies(rows, 0.3)
    empirical_rho07 = build_studies(rows, 0.7)
    template_indices = ten_study_template_indices(len(empirical_rho05))
    studies_rho05 = [empirical_rho05[index] for index in template_indices]
    studies_rho03 = [empirical_rho03[index] for index in template_indices]
    studies_rho07 = [empirical_rho07[index] for index in template_indices]
    x_gamma = [design_vector(study, "gamma") for study in studies_rho05]
    x_uniform = [design_vector(study, "uniform") for study in studies_rho05]
    x_lognormal = [design_vector(study, "lognormal") for study in studies_rho05]
    x_dmax_low = [design_vector(study, "gamma", 0.8) for study in studies_rho05]
    x_dmax_high = [design_vector(study, "gamma", 1.2) for study in studies_rho05]
    x_midpoint = [design_vector(study, "midpoint") for study in studies_rho05]

    result = {
        "replicates": replicates,
        "n_studies": len(studies_rho05),
        "n_comparisons": sum(len(study.rows) for study in studies_rho05),
        "n_point_comparisons": sum(
            row.refid == SCHULTZ_REFID
            for study in studies_rho05
            for row in study.rows
        ),
        "core_template": (
            "K=10: each of the eight empirical study-block configurations once plus "
            "two reproducibly resampled blocks"
        ),
        "core_template_indices_zero_based": template_indices,
        "parameter_uncertainty_design": (
            "Independent mean-preserving lognormal perturbations of each "
            "regional Gamma shape and scale parameter"
        ),
        "nonlinear_design": (
            "Individual-level logit includes curvature * log(1 + distance)"
        ),
        "inference_design": (
            "Model-based normal-Wald intervals compared with Student-t intervals "
            "using the number of independent study blocks minus one degree of freedom"
        ),
        "random_grid": [],
        "fixed_grid": [],
        "approach_comparison": [],
        "null_inference": [],
        "study_count_generalizability": [],
        "stress_tests": [],
    }

    multipliers = (0.5, 1.0, 2.0)
    for tau_mult in multipliers:
        for beta_mult in multipliers:
            beta = BASE_BETA * beta_mult
            tau = BASE_TAU * tau_mult
            seed = 1000 + int(beta_mult * 10) + 100 * int(tau_mult * 10)
            summary = simulate_linear(
                np.random.default_rng(seed),
                studies_rho05,
                studies_rho05,
                x_gamma,
                x_gamma,
                beta,
                tau,
                replicates,
            )
            result["random_grid"].append(summary)

            for true_name, true_x, approach_name, est_x in (
                ("Gamma", x_gamma, "Conditional mean", x_gamma),
                ("Gamma", x_gamma, "Midpoint", x_midpoint),
                ("Uniform", x_uniform, "Conditional mean", x_uniform),
                ("Uniform", x_uniform, "Midpoint", x_midpoint),
            ):
                # Use common random numbers for the two approaches within a
                # true-distribution cell.  Under a Uniform distribution the
                # conditional mean equals the midpoint, so this also provides
                # an exact implementation check.
                cell_seed = seed + sum(ord(c) for c in true_name)
                cell = simulate_linear(
                    np.random.default_rng(cell_seed),
                    studies_rho05,
                    studies_rho05,
                    true_x,
                    est_x,
                    beta,
                    tau,
                    replicates,
                )
                cell.update({"true_distribution": true_name, "approach": approach_name})
                result["approach_comparison"].append(cell)

    for tau_mult in multipliers:
        tau = BASE_TAU * tau_mult
        null_summary = simulate_linear(
            np.random.default_rng(6000 + 100 * int(tau_mult * 10)),
            studies_rho05,
            studies_rho05,
            x_gamma,
            x_gamma,
            0.0,
            tau,
            replicates,
        )
        null_summary["tau_multiplier"] = tau_mult
        result["null_inference"].append(null_summary)

    for n_studies in (10, 20, 30):
        alternative = simulate_resampled_study_count(
            np.random.default_rng(13000 + n_studies),
            empirical_rho05,
            [design_vector(study, "gamma") for study in empirical_rho05],
            BASE_BETA,
            BASE_TAU,
            replicates,
            n_studies,
        )
        alternative["effect_setting"] = "Empirical slope"
        result["study_count_generalizability"].append(alternative)

        null = simulate_resampled_study_count(
            np.random.default_rng(14000 + n_studies),
            empirical_rho05,
            [design_vector(study, "gamma") for study in empirical_rho05],
            0.0,
            BASE_TAU,
            replicates,
            n_studies,
        )
        null["effect_setting"] = "Null slope"
        result["study_count_generalizability"].append(null)

    for beta_mult in multipliers:
        beta = -0.011 * beta_mult
        result["fixed_grid"].append(
            simulate_fixed(
                np.random.default_rng(7000 + int(beta_mult * 10)),
                studies_rho05,
                x_gamma,
                beta,
                replicates,
            )
        )

    stress_specs = (
        ("Correct Gamma working model", studies_rho05, x_gamma, studies_rho05, x_gamma),
        ("Lognormal true; Gamma fitted", studies_rho05, x_lognormal, studies_rho05, x_gamma),
        ("Uniform true; Gamma fitted", studies_rho05, x_uniform, studies_rho05, x_gamma),
        ("True dmax 20% lower", studies_rho05, x_dmax_low, studies_rho05, x_gamma),
        ("True dmax 20% higher", studies_rho05, x_dmax_high, studies_rho05, x_gamma),
        ("True rho 0.3; fitted rho 0.5", studies_rho03, x_gamma, studies_rho05, x_gamma),
        ("True rho 0.7; fitted rho 0.5", studies_rho07, x_gamma, studies_rho05, x_gamma),
    )
    for index, (name, true_studies, true_x, est_studies, est_x) in enumerate(stress_specs):
        stress = simulate_linear(
            np.random.default_rng(9000 + index),
            true_studies,
            est_studies,
            true_x,
            est_x,
            BASE_BETA,
            BASE_TAU,
            replicates,
        )
        stress["scenario"] = name
        result["stress_tests"].append(stress)

    for index, beta_mult in enumerate(multipliers):
        beta = BASE_BETA * beta_mult
        stress = simulate_logistic_aggregation(
            np.random.default_rng(10000 + index),
            studies_rho05,
            x_gamma,
            beta,
            BASE_TAU,
            replicates,
        )
        stress["scenario"] = f"Individual-level logistic aggregation ({beta_mult:.1f}x beta)"
        result["stress_tests"].append(stress)

    for index, parameter_cv in enumerate((0.10, 0.20)):
        stress = simulate_parameter_uncertainty(
            np.random.default_rng(11000 + index),
            studies_rho05,
            x_gamma,
            BASE_BETA,
            BASE_TAU,
            replicates,
            parameter_cv,
        )
        stress["scenario"] = (
            "Gamma parameter uncertainty "
            f"(shape and scale CV {100 * parameter_cv:.0f}%)"
        )
        result["stress_tests"].append(stress)

    for index, curvature in enumerate((0.05, 0.10)):
        stress = simulate_logistic_aggregation(
            np.random.default_rng(12000 + index),
            studies_rho05,
            x_gamma,
            BASE_BETA,
            BASE_TAU,
            replicates,
            curvature=curvature,
        )
        stress["curvature"] = curvature
        stress["scenario"] = (
            "Nonlinear individual-level logit "
            f"(curvature {curvature:.2f} x log(1 + distance))"
        )
        result["stress_tests"].append(stress)

    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Reproduce the mixed-support simulation study."
    )
    parser.add_argument(
        "--data-file",
        type=Path,
        default=DEFAULT_BOOK,
        help=(
            "Path to the analytic Excel workbook (default: "
            "data/analytic_contrasts.xlsx)."
        ),
    )
    parser.add_argument("--replicates", type=int, default=5000)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()
    started = time.time()
    results = main(args.replicates, args.data_file.resolve())
    results["elapsed_seconds"] = time.time() - started
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(results, indent=2), encoding="utf-8")
    print(json.dumps(results, indent=2))
