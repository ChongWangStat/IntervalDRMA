"""Reproduce the residential-radon illustration from published summary data.

Pavia et al. (2003) reported adjusted odds ratios and confidence intervals for
study-specific intervals of time-weighted residential radon concentration.  No
exposure-category frequencies or outcome counts are used here.  Positive-width
supports are represented by conditional means under outcome-independent,
country-specific Lognormal distributions from UNSCEAR/WHO residential surveys.
"""

from __future__ import annotations

import argparse
import json
import math
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.stats import norm, t as student_t

from interval_drma_core import FastReml, Study, parametric_bootstrap_interval
from run_analysis import fit_fixed, residual_diagnostics


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_BOOK = ROOT / "data" / "radon_lung_cancer_contrasts.xlsx"
DEFAULT_OUTPUT = ROOT / "results" / "radon_example_results.json"

PAVIA_URL = "https://www.scielosp.org/pdf/bwho/2003.v81n10/732-738/en"
UNSCEAR_2000_URL = (
    "https://www.unscear.org/docs/publications/2000/UNSCEAR_2000_Annex-B.pdf"
)
WHO_2009_URL = "https://www.who.int/publications/i/item/9789241547673"
UNSCEAR_2024_URL = (
    "https://www.unscear.org/unscear/uploads/documents/unscear-reports/"
    "UNSCEAR_2024_Report_Vol.II.pdf"
)


def _records(sheet) -> list[dict]:
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value) for value in next(rows)]
    return [dict(zip(headers, values)) for values in rows if values[0] is not None]


def read_workbook(path: Path) -> tuple[list[dict], dict[str, dict[str, tuple[float, float]]]]:
    """Read source categories and historical/current distribution parameters."""
    workbook = load_workbook(path, data_only=False, read_only=True)
    categories = _records(workbook["Source Categories"])
    distribution_rows = _records(workbook["Country Distributions"])

    required_category = {
        "study_id", "study", "year", "country", "category_order",
        "reported_level", "lower_bq_m3", "upper_bq_m3", "open_upper",
        "reference", "adjusted_or", "ci_lower", "ci_upper",
    }
    if not categories or not required_category.issubset(categories[0]):
        raise ValueError("The source-category sheet is missing required columns")

    distributions: dict[str, dict[str, tuple[float, float]]] = {}
    for row in distribution_rows:
        country = str(row["country"])
        distributions[country] = {
            "historical": (float(row["primary_gm_bq_m3"]), float(row["primary_gsd"])),
            "2024": (float(row["sensitivity_gm_2024"]), float(row["sensitivity_gsd_2024"])),
        }
    return categories, distributions


def lognormal_conditional_mean(a: float, b: float, gm: float, gsd: float) -> float:
    """Return E[X | a < X <= b] for a Lognormal X; b may be infinite."""
    if not (0 <= a < b and gm > 0 and gsd > 1):
        raise ValueError("Invalid support or Lognormal parameters")
    mu = math.log(gm)
    sigma = math.log(gsd)
    upper_probability = 1.0 if math.isinf(b) else norm.cdf((math.log(b) - mu) / sigma)
    lower_probability = 0.0 if a == 0 else norm.cdf((math.log(a) - mu) / sigma)
    upper_first_moment = (
        1.0 if math.isinf(b)
        else norm.cdf((math.log(b) - mu - sigma**2) / sigma)
    )
    lower_first_moment = (
        0.0 if a == 0
        else norm.cdf((math.log(a) - mu - sigma**2) / sigma)
    )
    probability = upper_probability - lower_probability
    if probability <= 0:
        raise ValueError("Support has zero numerical probability")
    return math.exp(mu + sigma**2 / 2) * (
        upper_first_moment - lower_first_moment
    ) / probability


def build_contrasts(
    categories: list[dict],
    distributions: dict[str, dict[str, tuple[float, float]]],
    distribution_version: str = "historical",
    point_score_upper_multiplier: float | None = None,
) -> list[dict]:
    """Construct non-reference contrasts from support bounds and adjusted ORs."""
    contrasts = []
    study_ids = list(dict.fromkeys(str(row["study_id"]) for row in categories))
    for study_id in study_ids:
        block = [row for row in categories if str(row["study_id"]) == study_id]
        block.sort(key=lambda row: int(row["category_order"]))
        reference_rows = [row for row in block if bool(row["reference"])]
        if len(reference_rows) != 1:
            raise ValueError(f"{study_id} must have exactly one reference category")
        reference = reference_rows[0]
        country = str(reference["country"])
        gm, gsd = distributions[country][distribution_version]

        def representative(row: dict) -> float:
            lower = float(row["lower_bq_m3"])
            upper = math.inf if bool(row["open_upper"]) else float(row["upper_bq_m3"])
            if point_score_upper_multiplier is not None:
                return (
                    point_score_upper_multiplier * lower
                    if math.isinf(upper)
                    else (lower + upper) / 2
                )
            return lognormal_conditional_mean(lower, upper, gm, gsd)

        reference_mean = representative(reference)
        for row in block:
            if bool(row["reference"]):
                continue
            odds_ratio = float(row["adjusted_or"])
            ci_lower = float(row["ci_lower"])
            ci_upper = float(row["ci_upper"])
            if not (odds_ratio > 0 and 0 < ci_lower < ci_upper):
                raise ValueError(f"Invalid OR or CI in {study_id}")
            comparison_mean = representative(row)
            contrasts.append({
                "study_id": study_id,
                "country": country,
                "comparison_level": str(row["reported_level"]),
                "reference_level": str(reference["reported_level"]),
                "log_or": math.log(odds_ratio),
                "se_log_or": (math.log(ci_upper) - math.log(ci_lower)) / 3.92,
                "comparison_mean": comparison_mean,
                "reference_mean": reference_mean,
                "x": comparison_mean - reference_mean,
            })
    return contrasts


def build_analysis_data(contrasts: list[dict], rho: float):
    studies, design, outcomes = [], [], []
    study_ids = list(dict.fromkeys(row["study_id"] for row in contrasts))
    for index, study_id in enumerate(study_ids, start=1):
        block = [row for row in contrasts if row["study_id"] == study_id]
        standard_errors = np.asarray([row["se_log_or"] for row in block], dtype=float)
        correlation = np.full((len(block), len(block)), rho, dtype=float)
        np.fill_diagonal(correlation, 1.0)
        covariance = np.diag(standard_errors) @ correlation @ np.diag(standard_errors)
        studies.append(Study(index, [], covariance))
        design.append(np.asarray([row["x"] for row in block], dtype=float))
        outcomes.append(np.asarray([row["log_or"] for row in block], dtype=float))
    return studies, design, outcomes


def cr2_slope_sensitivity(studies, design, outcomes, beta: float, tau: float) -> dict:
    """Return a one-parameter CR2 sandwich sensitivity at the fitted tau.

    Each marginal covariance block is whitened before applying the CR2
    leverage correction.  A conservative t critical value with K-1 degrees of
    freedom is used; this is a sensitivity analysis rather than the primary
    model-based interval.
    """
    whitened_design = []
    whitened_outcomes = []
    for study, xk, yk in zip(studies, design, outcomes):
        marginal = study.sigma + tau**2 * np.outer(xk, xk)
        factor = np.linalg.cholesky(marginal)
        whitened_design.append(np.linalg.solve(factor, xk))
        whitened_outcomes.append(np.linalg.solve(factor, yk))

    information = float(sum(xk @ xk for xk in whitened_design))
    meat = 0.0
    minimum_adjustment_eigenvalue = math.inf
    for xk, yk in zip(whitened_design, whitened_outcomes):
        residual = yk - beta * xk
        complement = np.eye(len(xk)) - np.outer(xk, xk) / information
        eigenvalues, eigenvectors = np.linalg.eigh(
            0.5 * (complement + complement.T)
        )
        minimum_adjustment_eigenvalue = min(
            minimum_adjustment_eigenvalue, float(eigenvalues.min())
        )
        if eigenvalues.min() <= 1e-10:
            raise RuntimeError("CR2 leverage adjustment is numerically singular")
        adjustment = (
            eigenvectors
            @ np.diag(1.0 / np.sqrt(eigenvalues))
            @ eigenvectors.T
        )
        score = float(xk @ adjustment @ residual)
        meat += score**2

    se = math.sqrt(meat) / information
    degrees_freedom = len(studies) - 1
    critical = float(student_t.ppf(0.975, degrees_freedom))
    return {
        "method": (
            "CR2 cluster-sandwich sensitivity after whitening by the fitted "
            "marginal covariance; t critical value with K-1 degrees of freedom"
        ),
        "beta": float(beta),
        "se_beta": float(se),
        "df": degrees_freedom,
        "beta_ci": [float(beta - critical * se), float(beta + critical * se)],
        "minimum_leverage_complement_eigenvalue": minimum_adjustment_eigenvalue,
    }


def fit_contrasts(
    contrasts: list[dict],
    rho: float,
    bootstrap_replicates: int = 0,
    bootstrap_seed: int = 161803,
) -> dict:
    studies, design, outcomes = build_analysis_data(contrasts, rho)
    random_effects = FastReml(studies, design).fit_detailed(outcomes)
    random_effects["cr2_sensitivity"] = cr2_slope_sensitivity(
        studies,
        design,
        outcomes,
        random_effects["beta"],
        random_effects["tau"],
    )
    marginal_information = []
    condition_numbers = []
    for study, xk in zip(studies, design):
        marginal = study.sigma + random_effects["tau"] ** 2 * np.outer(xk, xk)
        marginal_information.append(float(xk @ np.linalg.solve(marginal, xk)))
        condition_numbers.append(
            {
                "study_refid": int(study.refid),
                "sampling_covariance_condition_number": float(
                    np.linalg.cond(study.sigma)
                ),
                "marginal_covariance_condition_number": float(np.linalg.cond(marginal)),
            }
        )
    total_information = sum(marginal_information)
    random_effects["study_information_weights"] = [
        {
            "study_refid": int(study.refid),
            "weight": float(value / total_information),
        }
        for study, value in zip(studies, marginal_information)
    ]
    random_effects["covariance_condition_numbers"] = condition_numbers
    if bootstrap_replicates:
        random_effects["parametric_bootstrap"] = parametric_bootstrap_interval(
            studies,
            design,
            outcomes,
            replicates=bootstrap_replicates,
            seed=bootstrap_seed,
            fitted=random_effects,
        )
    random_effects["or_per_100_bq_m3"] = math.exp(100 * random_effects["beta"])
    random_effects["or_ci_per_100_bq_m3"] = [
        math.exp(100 * value) for value in random_effects["beta_ci"]
    ]
    random_effects["or_at_150_bq_m3"] = math.exp(150 * random_effects["beta"])
    random_effects["or_ci_at_150_bq_m3"] = [
        math.exp(150 * value) for value in random_effects["beta_ci"]
    ]
    random_effects["residual_diagnostics"] = residual_diagnostics(
        studies, design, outcomes, random_effects["beta"], random_effects["tau"]
    )
    fixed_effects = fit_fixed(studies, design, outcomes)
    fixed_effects["or_per_100_bq_m3"] = math.exp(100 * fixed_effects["beta"])
    fixed_effects["or_ci_per_100_bq_m3"] = [
        math.exp(100 * value) for value in fixed_effects["beta_ci"]
    ]
    return {
        "rho": rho,
        "n_studies": len(studies),
        "n_contrasts": len(contrasts),
        "random_effects": random_effects,
        "fixed_effects": fixed_effects,
    }


def leave_one_study_out(contrasts: list[dict], rho: float = 0.5) -> list[dict]:
    """Refit the primary model after omitting each published study."""
    study_ids = list(dict.fromkeys(row["study_id"] for row in contrasts))
    results = []
    for study_id in study_ids:
        retained = [row for row in contrasts if row["study_id"] != study_id]
        fitted = fit_contrasts(retained, rho=rho)
        random_effects = fitted["random_effects"]
        results.append(
            {
                "omitted_study": study_id,
                "n_studies": fitted["n_studies"],
                "n_contrasts": fitted["n_contrasts"],
                "beta": random_effects["beta"],
                "se_beta": random_effects["se_beta"],
                "beta_ci": random_effects["beta_ci"],
                "tau": random_effects["tau"],
                "or_per_100_bq_m3": random_effects["or_per_100_bq_m3"],
                "or_ci_per_100_bq_m3": random_effects["or_ci_per_100_bq_m3"],
            }
        )
    return results


def run_analysis(path: Path, bootstrap_replicates: int = 2000) -> dict:
    categories, distributions = read_workbook(path)
    primary_contrasts = build_contrasts(categories, distributions, "historical")
    primary = fit_contrasts(
        primary_contrasts,
        rho=0.5,
        bootstrap_replicates=bootstrap_replicates,
    )
    representative_contrasts = []
    for study_id in dict.fromkeys(row["study_id"] for row in primary_contrasts):
        candidates = [
            row for row in primary_contrasts if row["study_id"] == study_id
        ]
        representative_contrasts.append(
            min(candidates, key=lambda row: row["se_log_or"])
        )
    representative_sensitivity = fit_contrasts(
        representative_contrasts,
        rho=0.0,
    )
    correlation_sensitivity = [
        fit_contrasts(primary_contrasts, rho=rho)
        for rho in (0.0, 0.25, 0.5, 0.75, 0.9)
        if rho != 0.5
    ]
    current_contrasts = build_contrasts(categories, distributions, "2024")
    distribution_sensitivity = fit_contrasts(current_contrasts, rho=0.5)
    leave_one_out = leave_one_study_out(primary_contrasts)
    point_score_sensitivity = []
    for multiplier in (1.2, 1.5, 2.0):
        point_contrasts = build_contrasts(
            categories,
            distributions,
            "historical",
            point_score_upper_multiplier=multiplier,
        )
        fitted = fit_contrasts(point_contrasts, rho=0.5)
        fitted["bounded_score"] = "midpoint"
        fitted["right_open_score"] = f"{multiplier:.1f} times lower bound"
        point_score_sensitivity.append(fitted)

    return {
        "data_file": path.name,
        "outcome": "lung cancer",
        "exposure": "time-weighted residential radon concentration in Bq/m^3",
        "estimand": "change in log odds ratio per Bq/m^3 higher representative radon concentration",
        "source_meta_analysis": PAVIA_URL,
        "reml_criterion_note": (
            "The profiled REML criterion is an optimization diagnostic within a "
            "fixed design mapping and is not comparable across mappings that "
            "change the fixed-effects design vector."
        ),
        "distribution_sources": [UNSCEAR_2000_URL, WHO_2009_URL, UNSCEAR_2024_URL],
        "excluded_study": (
            "Letourneau et al. (1994), because Pavia et al. report only "
            "cumulative exposure rather than time-weighted Bq/m^3"
        ),
        "first_stage": (
            "Country-specific outcome-independent Lognormal working distributions. "
            "Historical survey parameters are primary and UNSCEAR 2024 parameters "
            "are a sensitivity analysis. Integer-rounded categories use contiguous "
            "half-unit cutpoints; decimal cutpoints are retained as reported. No "
            "exposure-category frequencies are used."
        ),
        "n_studies": len(set(row["study_id"] for row in primary_contrasts)),
        "n_contrasts": len(primary_contrasts),
        "primary_historical_distribution": primary,
        "one_most_precise_contrast_per_study_sensitivity": {
            "selection_rule": (
                "One adjusted OR contrast with the smallest log-scale standard "
                "error was retained per study; rho=0 because one contrast remains per block"
            ),
            **representative_sensitivity,
        },
        "correlation_sensitivity": correlation_sensitivity,
        "unscear_2024_distribution_sensitivity": distribution_sensitivity,
        "leave_one_study_out": leave_one_out,
        "point_score_sensitivity": point_score_sensitivity,
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Reproduce the radon illustration")
    parser.add_argument("--data-file", type=Path, default=DEFAULT_BOOK)
    parser.add_argument("--bootstrap-replicates", type=int, default=2000)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    arguments = parser.parse_args()
    results = run_analysis(
        arguments.data_file.resolve(),
        bootstrap_replicates=arguments.bootstrap_replicates,
    )
    arguments.output.parent.mkdir(parents=True, exist_ok=True)
    arguments.output.write_text(json.dumps(results, indent=2), encoding="utf-8")
    print(json.dumps(results, indent=2))
